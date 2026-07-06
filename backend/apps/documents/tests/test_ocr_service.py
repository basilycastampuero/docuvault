import io
from unittest.mock import MagicMock, patch

import pytest
from botocore.exceptions import ClientError

from apps.audit.models import AuditAction, AuditLog
from apps.core.exceptions import TransientError
from apps.documents.models import OcrStatus
from apps.documents.services import ocr_service
from apps.documents.tests.factories import DocumentFactory
from apps.search.selectors.search_selector import search_documents

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _client_error(code: str) -> ClientError:
    return ClientError({"Error": {"Code": code, "Message": ""}}, "GetObject")


def _real_docx_bytes(paragraphs: list[str]) -> bytes:
    """Build a genuine DOCX file (real python-docx round trip, no mocking)."""
    from docx import Document as DocxDocument

    doc = DocxDocument()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _real_xlsx_bytes(sheets: dict[str, list[list]]) -> bytes:
    """Build a genuine XLSX file (real openpyxl round trip, no mocking)."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(sheet_name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def mock_storage(monkeypatch):
    """Patch StorageService in ocr_service; return the instance mock."""
    instance = MagicMock()
    instance.download_file.return_value = b"blob-bytes"
    cls = MagicMock(return_value=instance)
    monkeypatch.setattr("apps.documents.services.ocr_service.StorageService", cls)
    return instance


@pytest.mark.django_db
class TestOcrServiceProcess:
    def test_image_is_ocrd_and_completed(self, mock_storage):
        doc = DocumentFactory(mime_type="image/png", ocr_content="")
        with (
            patch("apps.documents.services.ocr_service.Image.open"),
            patch(
                "apps.documents.services.ocr_service.pytesseract.image_to_string",
                return_value="hello from image",
            ),
        ):
            ocr_service.process(doc)
        doc.refresh_from_db()
        assert doc.ocr_status == OcrStatus.COMPLETED
        assert doc.ocr_content == "hello from image"

    def test_pdf_is_ocrd_page_by_page(self, mock_storage):
        doc = DocumentFactory(mime_type="application/pdf", ocr_content="")
        with (
            patch(
                "apps.documents.services.ocr_service.convert_from_bytes",
                return_value=["page1", "page2"],
            ),
            patch(
                "apps.documents.services.ocr_service.pytesseract.image_to_string",
                side_effect=["text one", "text two"],
            ),
        ):
            ocr_service.process(doc)
        doc.refresh_from_db()
        assert doc.ocr_status == OcrStatus.COMPLETED
        assert doc.ocr_content == "text one\ntext two"

    def test_unsupported_mime_is_skipped(self, mock_storage):
        # Legacy .doc (not OOXML) remains unsupported even after Phase 6.2 added
        # DOCX/XLSX extraction — only true OOXML mimes are handled.
        doc = DocumentFactory(mime_type="application/msword")
        ocr_service.process(doc)
        doc.refresh_from_db()
        assert doc.ocr_status == OcrStatus.SKIPPED
        mock_storage.download_file.assert_not_called()

    def test_missing_blob_marks_failed_without_retry(self, mock_storage):
        mock_storage.download_file.side_effect = _client_error("NoSuchKey")
        doc = DocumentFactory(mime_type="image/png")
        ocr_service.process(doc)  # must not raise
        doc.refresh_from_db()
        assert doc.ocr_status == OcrStatus.FAILED

    def test_storage_timeout_raises_transient_error(self, mock_storage):
        mock_storage.download_file.side_effect = _client_error("500")
        doc = DocumentFactory(mime_type="image/png")
        with pytest.raises(TransientError):
            ocr_service.process(doc)

    def test_network_error_raises_transient_error(self, mock_storage):
        mock_storage.download_file.side_effect = ConnectionError("timeout")
        doc = DocumentFactory(mime_type="image/png")
        with pytest.raises(TransientError):
            ocr_service.process(doc)

    def test_corrupt_file_marks_failed_without_retry(self, mock_storage):
        doc = DocumentFactory(mime_type="image/png")
        with patch(
            "apps.documents.services.ocr_service.Image.open",
            side_effect=OSError("cannot identify image file"),
        ):
            ocr_service.process(doc)  # must not raise
        doc.refresh_from_db()
        assert doc.ocr_status == OcrStatus.FAILED

    def test_blank_page_completes_with_empty_content(self, mock_storage):
        doc = DocumentFactory(mime_type="image/png", ocr_content="old")
        with (
            patch("apps.documents.services.ocr_service.Image.open"),
            patch(
                "apps.documents.services.ocr_service.pytesseract.image_to_string",
                return_value="   ",
            ),
        ):
            ocr_service.process(doc)
        doc.refresh_from_db()
        assert doc.ocr_status == OcrStatus.COMPLETED
        assert doc.ocr_content == ""

    def test_completion_is_audited_with_via_ocr(self, mock_storage):
        doc = DocumentFactory(mime_type="image/png")
        with (
            patch("apps.documents.services.ocr_service.Image.open"),
            patch(
                "apps.documents.services.ocr_service.pytesseract.image_to_string",
                return_value="text",
            ),
        ):
            ocr_service.process(doc)
        log = AuditLog.objects.filter(
            entity_id=str(doc.id), action=AuditAction.UPDATE
        ).latest("created_at")
        assert log.metadata == {"via": "ocr"}
        assert log.user is None  # system action

    def test_document_is_searchable_by_ocr_content(self, mock_storage):
        """The DoD: after OCR, the document is found by a word from its content."""
        doc = DocumentFactory(
            organization__name="Acme", mime_type="image/png", ocr_content=""
        )
        with (
            patch("apps.documents.services.ocr_service.Image.open"),
            patch(
                "apps.documents.services.ocr_service.pytesseract.image_to_string",
                return_value="contrato de arrendamiento confidencial",
            ),
        ):
            ocr_service.process(doc)
        results = search_documents(organization=doc.organization, query="arrendamiento")
        assert doc in list(results)


@pytest.mark.django_db
class TestOcrServiceOfficeExtraction:
    def test_docx_text_is_extracted_and_completed(self, mock_storage):
        """Real DOCX (python-docx round trip): paragraph text ends up in ocr_content."""
        mock_storage.download_file.return_value = _real_docx_bytes(
            ["Contrato de arrendamiento", "Cláusula segunda: confidencialidad"]
        )
        doc = DocumentFactory(mime_type=_DOCX_MIME, ocr_content="")

        ocr_service.process(doc)

        doc.refresh_from_db()
        assert doc.ocr_status == OcrStatus.COMPLETED
        assert "Contrato de arrendamiento" in doc.ocr_content
        assert "confidencialidad" in doc.ocr_content

    def test_xlsx_cells_are_extracted_and_flattened(self, mock_storage):
        """Real XLSX (openpyxl round trip): every cell value ends up in ocr_content."""
        mock_storage.download_file.return_value = _real_xlsx_bytes(
            {
                "Sheet1": [["Cliente", "Monto"], ["Acme Corp", 1500]],
                "Sheet2": [["Segunda hoja", "valor"]],
            }
        )
        doc = DocumentFactory(mime_type=_XLSX_MIME, ocr_content="")

        ocr_service.process(doc)

        doc.refresh_from_db()
        assert doc.ocr_status == OcrStatus.COMPLETED
        assert "Cliente" in doc.ocr_content
        assert "Acme Corp" in doc.ocr_content
        assert "1500" in doc.ocr_content
        assert "Segunda hoja" in doc.ocr_content

    def test_docx_document_is_searchable_by_extracted_content(self, mock_storage):
        """End-to-end FTS: a word from the DOCX text is found via the search selector."""
        mock_storage.download_file.return_value = _real_docx_bytes(
            ["Presupuesto anual confidencial 2026"]
        )
        doc = DocumentFactory(mime_type=_DOCX_MIME, ocr_content="")

        ocr_service.process(doc)

        results = search_documents(organization=doc.organization, query="presupuesto")
        assert doc in list(results)

    def test_xlsx_document_is_searchable_by_extracted_content(self, mock_storage):
        """End-to-end FTS: a cell value from the XLSX is found via the search selector."""
        mock_storage.download_file.return_value = _real_xlsx_bytes(
            {"Sheet1": [["Factura", "Proveedor Contoso"]]}
        )
        doc = DocumentFactory(mime_type=_XLSX_MIME, ocr_content="")

        ocr_service.process(doc)

        results = search_documents(organization=doc.organization, query="contoso")
        assert doc in list(results)

    @pytest.mark.parametrize(
        "mime_type",
        ["application/vnd.ms-excel", "application/zip"],
    )
    def test_other_legacy_or_archive_mimes_are_skipped(self, mock_storage, mime_type):
        """Legacy .xls and plain .zip must never reach the docx/xlsx handlers."""
        doc = DocumentFactory(mime_type=mime_type)

        ocr_service.process(doc)

        doc.refresh_from_db()
        assert doc.ocr_status == OcrStatus.SKIPPED
        mock_storage.download_file.assert_not_called()

    def test_corrupt_docx_marks_failed_without_retry(self, mock_storage):
        mock_storage.download_file.return_value = b"this is not a real zip/docx file"
        doc = DocumentFactory(mime_type=_DOCX_MIME)

        ocr_service.process(doc)  # must not raise

        doc.refresh_from_db()
        assert doc.ocr_status == OcrStatus.FAILED

    def test_corrupt_xlsx_marks_failed_without_retry(self, mock_storage):
        mock_storage.download_file.return_value = b"this is not a real zip/xlsx file"
        doc = DocumentFactory(mime_type=_XLSX_MIME)

        ocr_service.process(doc)  # must not raise

        doc.refresh_from_db()
        assert doc.ocr_status == OcrStatus.FAILED
